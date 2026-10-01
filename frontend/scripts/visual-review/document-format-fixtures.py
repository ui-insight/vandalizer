"""Build bounded local viewer inputs, using the application's actual format readers.

Run from the repo root with backend/.venv/bin/python. No database/model/network calls.
Outputs are disposable QA inputs, not user documents or production uploads.
"""
import json
import sys
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

repo = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(repo / 'backend'))
from app.services.document_readers import read_docx_markdown, extract_sheet_json_from_xlsx

out = Path(sys.argv[1] if len(sys.argv) > 1 else '/private/tmp/vandalizer-ra8-documents')
out.mkdir(parents=True, exist_ok=True)
source = fitz.open(repo / 'backend/certification-data/documents/budget-justification.pdf')
scan = canvas.Canvas(str(out / 'scanned-budget.pdf'))
for page in source:
    image = out / f'scan-source-{page.number}.png'
    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(image)
    scan.setPageSize((page.rect.width, page.rect.height))
    scan.drawImage(ImageReader(str(image)), 0, 0, width=page.rect.width, height=page.rect.height)
    scan.showPage()
scan.save()
assert all(not page.get_text().strip() for page in fitz.open(out / 'scanned-budget.pdf'))
(out / 'pdf-text.json').write_text(json.dumps({'text': '\n\n'.join(page.get_text() for page in source)}))

doc = Document()
doc.add_heading('Synthetic sponsor budget review', 0)
doc.add_paragraph('For UI QA only. Submission deadline: October 15, 2026 at 5:00 PM Pacific. Confirm exception evidence before approval.')
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for cell, value in zip(table.rows[0].cells, ['Category', 'Requested', 'Evidence']):
    cell.text = value
for values in [('Senior personnel', '$43,333', 'Salary justification required'), ('Equipment', '$45,000', 'Quote and sponsor exception'), ('Travel', '$40,000', 'Conference and field visits')]:
    for cell, value in zip(table.add_row().cells, values):
        cell.text = value
doc.add_paragraph('Known limitation: this is a small synthetic fixture, not a sponsor rule.')
doc.save(out / 'sponsor-review.docx')
text = read_docx_markdown(str(out / 'sponsor-review.docx'))
assert '43,333' in text and 'Equipment' in text
(out / 'word-text.json').write_text(json.dumps({'text': text}))

book = Workbook()
budget = book.active
budget.title = 'Budget'
for row in [['Category', 'Requested', 'Institutional contribution', 'Evidence'], ['Senior personnel', 43333, 5000, 'Salary justification required'], ['Equipment', 45000, 0, 'Sponsor exception'], ['Travel', 40000, 2000, 'Conference and field visits'], ['Total', '=SUM(B2:B4)', '=SUM(C2:C4)', 'Confirm budget justification']]:
    budget.append(row)
assumptions = book.create_sheet('Assumptions')
assumptions.append(['Item', 'Explanation'])
assumptions.append(['Deadline', 'October 15, 2026 at 5 PM Pacific'])
book.save(out / 'budget.xlsx')
sheets = extract_sheet_json_from_xlsx(str(out / 'budget.xlsx'))
assert float(sheets['sheets'][0]['rows'][-1][1]) == 128333
(out / 'sheet-data.json').write_text(json.dumps(sheets))
print(json.dumps({'pdf_pages': len(source), 'word_text_characters': len(text), 'worksheets': len(sheets['sheets']), 'evaluated_total': sheets['sheets'][0]['rows'][-1][1]}))
