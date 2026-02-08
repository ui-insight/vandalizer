#!/usr/bin/env python3
"""Output handlers for passive workflows: storage, notifications, webhooks, chains."""

import csv
import json
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from flask import render_template
from flask_mail import Message

from app import app, mail
from app.models import WorkflowResult, SmartDocument, SmartFolder, Workflow, WorkItem


def save_results_to_folder(result, storage_config):
    """
    Save workflow results to a folder in the configured format.
    
    Args:
        result: WorkflowResult instance
        storage_config: Dict with destination_folder, file_naming, format, etc.
        
    Returns:
        String path to the saved file
    """
    
    folder_id = storage_config.get("destination_folder")
    if not folder_id:
        raise ValueError("No destination folder configured")
    
    folder = SmartFolder.objects(uuid=folder_id).first()
    if not folder:
        raise ValueError(f"Folder {folder_id} not found")
    
    # Generate filename from template
    filename_template = storage_config.get("file_naming", "{date}_{workflow_name}_results")
    filename = filename_template.format(
        date=datetime.now().strftime("%Y-%m-%d"),
        time=datetime.now().strftime("%H-%M-%S"),
        workflow_name=result.workflow.name.replace(" ", "_"),
        workflow_id=str(result.workflow.id),
        run_id=str(result.id)
    )
    
    # Get format
    format_type = storage_config.get("format", "csv")
    filename = f"{filename}.{format_type}"
    
    # Get results data
    final_output = result.final_output or {}
    output_data = final_output.get("output")
    
    # Create upload directory
    upload_dir = Path(app.root_path) / "static" / "uploads" / result.workflow.user_id
    upload_dir.mkdir(parents=True, exist_ok=True)
    
    file_path = upload_dir / filename
    
    # Save based on format
    if format_type == "csv":
        save_as_csv(file_path, output_data)
    elif format_type == "json":
        save_as_json(file_path, output_data)
    elif format_type in ["xlsx", "excel"]:
        save_as_excel(file_path, output_data)
    else:
        # Default to text
        with open(file_path, "w") as f:
            f.write(str(output_data))
    
    # Create SmartDocument for the output file
    doc_uuid = uuid4().hex
    
    doc = SmartDocument(
        title=filename,
        path=f"{result.workflow.user_id}/{filename}",
        downloadpath=f"{result.workflow.user_id}/{filename}",
        extension=format_type,
        uuid=doc_uuid,
        user_id=result.workflow.user_id,
        space=folder.space,
        folder=folder.uuid
    )
    doc.save()
    
    return str(file_path)


def save_as_csv(file_path, data):
    """Save data as CSV."""
    if not data:
        with open(file_path, "w") as f:
            f.write("")
        return
    
    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
        # List of dicts -> CSV with headers
        with open(file_path, "w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=data[0].keys())
            writer.writeheader()
            writer.writerows(data)
    else:
        # Plain text
        with open(file_path, "w") as f:
            f.write(str(data))


def save_as_json(file_path, data):
    """Save data as JSON."""
    with open(file_path, "w") as f:
        json.dump(data, f, indent=2, default=str)


def save_as_excel(file_path, data):
    """Save data as Excel (requires openpyxl)."""
    try:
        import openpyxl
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        
        if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
            # Write headers
            headers = list(data[0].keys())
            ws.append(headers)
            
            # Write rows
            for row in data:
                ws.append([row.get(h) for h in headers])
        else:
            # Write as single cell
            ws["A1"] = str(data)
        
        wb.save(file_path)
    except ImportError:
        # Fallback to JSON if openpyxl not available
        save_as_json(file_path.with_suffix(".json"), data)


def should_send_notification(result, notification):
    """
    Check if notification should be sent based on conditions.
    
    Args:
        result: WorkflowResult instance
        notification: Dict with conditions key
        
    Returns:
        Boolean indicating if notification should be sent
    """
    
    condition = notification.get("conditions", "always")
    
    if condition == "always":
        return True
    elif condition == "success":
        return result.status == "completed"
    elif condition == "failure":
        return result.status == "failed"
    
    return True


def send_workflow_notification(result, notification, work_item=None):
    """
    Send notification for workflow completion.

    Args:
        result: WorkflowResult instance
        notification: Dict with channel, recipients, settings
        work_item: Optional WorkItem instance (for M365 Teams notifications)
    """

    channel = notification.get("channel", "email")
    if channel == "teams":
        _send_teams_notification(result, notification, work_item)
        return
    if channel != "email":
        return
    
    # Build recipient list
    recipients = list(notification.get("recipients", []))
    
    if notification.get("notify_owner"):
        workflow = result.workflow
        from app.models import User
        user = User.objects(user_id=workflow.user_id).first()
        if user and user.email and user.email not in recipients:
            recipients.append(user.email)
    
    if notification.get("notify_team"):
        # TODO: Get team members' emails
        pass
    
    if not recipients:
        return
    
    # Determine subject
    if result.status == "completed":
        subject = f"✓ {result.workflow.name} completed"
    elif result.status == "failed":
        subject = f"⚠️ {result.workflow.name} failed"
    else:
        subject = f"{result.workflow.name} - {result.status}"
    
    # Render email template
    try:
        html_body = render_template(
            "emails/workflow_notification.html",
            workflow=result.workflow,
            result=result,
            event=result.trigger_event,
            include_summary=notification.get("include_summary", True),
            include_results=notification.get("include_full_results", False)
        )
    except Exception:
        # Fallback to simple text if template not found
        html_body = f"""
        <html>
        <body>
            <h2>{subject}</h2>
            <p>Your workflow "{result.workflow.name}" has {result.status}.</p>
            <p><strong>Trigger Type:</strong> {result.trigger_type or 'manual'}</p>
            <p><strong>Started:</strong> {result.start_time}</p>
            <p><strong>Duration:</strong> {(result.trigger_event.duration_ms / 1000) if result.trigger_event and result.trigger_event.duration_ms else 'N/A'} seconds</p>
        </body>
        </html>
        """
    
    # Send email
    msg = Message(
        subject=subject,
        recipients=recipients,
        html=html_body,
        sender=app.config.get("MAIL_DEFAULT_SENDER", "noreply@vandalizer.com")
    )
    
    mail.send(msg)


def save_results_to_onedrive_channel(result, onedrive_config, work_item=None):
    """Save workflow results to a OneDrive case folder.

    Args:
        result: WorkflowResult instance.
        onedrive_config: Dict with drive_id, base_path, etc.
        work_item: Optional WorkItem instance. If None, tries to find one
            via the trigger event.

    Returns:
        The case folder path string.
    """
    if not work_item:
        event = result.trigger_event
        if event:
            work_item = WorkItem.objects(trigger_event=event).first()
    if not work_item:
        raise ValueError("No WorkItem available for OneDrive output")

    from app.utilities.onedrive_output import save_results_to_onedrive

    return save_results_to_onedrive(result, work_item, onedrive_config)


def _send_teams_notification(result, notification, work_item=None):
    """Send a Teams Adaptive Card notification for a workflow result.

    Args:
        result: WorkflowResult instance.
        notification: Dict with team_id, channel_id settings.
        work_item: Optional WorkItem for richer card content.
    """
    team_id = notification.get("team_id")
    channel_id = notification.get("channel_id")
    if not team_id or not channel_id:
        return

    # Resolve work item if not provided
    if not work_item:
        event = result.trigger_event
        if event:
            work_item = WorkItem.objects(trigger_event=event).first()

    # Determine which card to build
    from app.utilities.teams_cards import build_work_item_card, build_exception_card

    if result.status == "failed":
        error = str(result.final_output.get("error", "Unknown error")) if result.final_output else "Unknown error"
        if work_item:
            card = build_exception_card(work_item, error)
        else:
            return  # Can't build a useful card without a work item
    elif work_item:
        card = build_work_item_card(work_item, result=result)
    else:
        return

    # Find a user_id to authenticate the Graph call
    user_id = None
    if work_item:
        user_id = work_item.owner_user_id
    if not user_id:
        user_id = result.workflow.user_id if result.workflow else None
    if not user_id:
        return

    try:
        from app.utilities.graph_client import GraphClient

        client = GraphClient(user_id)
        client.send_channel_message(team_id, channel_id, card_json=card)
    except Exception:
        import logging
        logging.getLogger(__name__).warning(
            f"Failed to send Teams notification for result {result.id}"
        )


def call_webhook(result, webhook_config):
    """
    Call a webhook with workflow result data.
    
    Args:
        result: WorkflowResult instance
        webhook_config: Dict with url, method, auth, payload_template
        
    Returns:
        Response object
    """
    import requests
    
    url = webhook_config.get("url")
    method = webhook_config.get("method", "POST").upper()
    
    # Build payload
    payload_template = webhook_config.get("payload_template")
    if payload_template:
        # Template substitution
        payload = payload_template.format(
            workflow_id=str(result.workflow.id),
            workflow_name=result.workflow.name,
            result_id=str(result.id),
            status=result.status,
            output=json.dumps(result.final_output.get("output") if result.final_output else None)
        )
        payload = json.loads(payload)
    else:
        # Default payload
        payload = {
            "workflow_id": str(result.workflow.id),
            "workflow_name": result.workflow.name,
            "result_id": str(result.id),
            "status": result.status,
            "trigger_type": result.trigger_type,
            "started_at": result.start_time.isoformat() if result.start_time else None,
            "output": result.final_output.get("output") if result.final_output else None
        }
    
    # Setup authentication
    headers = {"Content-Type": "application/json"}
    auth_config = webhook_config.get("auth", {})
    auth_type = auth_config.get("type")
    
    if auth_type == "bearer":
        headers["Authorization"] = f"Bearer {auth_config.get('token')}"
    elif auth_type == "api_key":
        key_name = auth_config.get("key_name", "X-API-Key")
        headers[key_name] = auth_config.get("api_key")
    
    # Make request
    if method == "POST":
        response = requests.post(url, json=payload, headers=headers, timeout=30)
    elif method == "PUT":
        response = requests.put(url, json=payload, headers=headers, timeout=30)
    else:
        raise ValueError(f"Unsupported HTTP method: {method}")
    
    response.raise_for_status()
    return response
