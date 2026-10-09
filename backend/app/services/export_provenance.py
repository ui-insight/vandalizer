"""Keep the evidence when a result leaves Vandalizer (#1008).

Extraction captures a supporting quote and page per field, and workflow
knowledge-base steps record the passages they retrieved, but no download
carried either: a colleague or PI handed the export had no way to check it.
These helpers collect a run's sources and render them as a Sources section
(Markdown / Word / PDF) or as per-field columns (CSV), and check the quotes a
chat-saved document cites against the documents they name.
"""

import re
from typing import Any

SOURCE_HEADERS = ["Item", "Document", "Page", "Quote", "Status"]

_SUPPORT_LABELS = {  # extraction_sources.SUPPORT_* states
    "supported": "quote found; value matches it",
    "unassessed": "quote found",
    "quote_unsupported": "quote found; value not in it",
    "unverified": "quote not found in the document",
}


def _page_label(entry: dict) -> str:
    page = entry.get("page")
    if not isinstance(page, int):
        return ""
    end = entry.get("page_end")
    approx = "~" if entry.get("page_approximate") else ""
    if isinstance(end, int) and end > page:
        return f"p. {approx}{page}–{end}"
    return f"p. {approx}{page}"


def _field_status(entry: dict) -> str:
    support = entry.get("support")
    if isinstance(support, str) and support in _SUPPORT_LABELS:
        return _SUPPORT_LABELS[support]
    if entry.get("verified") is True:
        return "quote found"
    if entry.get("verified") is False:
        return "quote not found in the document"
    return ""


def _step_field_sources(step: Any) -> list[dict]:
    if isinstance(step, dict) and isinstance(step.get("field_sources"), list):
        return [s for s in step["field_sources"] if isinstance(s, dict)]
    return []


def output_step_field_sources(status: dict) -> list[dict]:
    """``field_sources`` of the step whose output a single-file export holds.

    Positional against that step's output: index *i* maps field names to the
    source of output row *i* (see ``MultiTaskNode``).
    """
    steps = status.get("steps_output") or {}
    names = [n for n in (status.get("output_step_names") or []) if n in steps]
    if len(names) == 1:
        return _step_field_sources(steps[names[0]])
    if steps:
        return _step_field_sources(list(steps.values())[-1])
    return []


def collect_run_sources(status: dict) -> list[dict]:
    """Every source a workflow run recorded, as rows for a Sources section."""
    rows: list[dict] = []
    seen: set[tuple] = set()
    for step in (status.get("steps_output") or {}).values():
        for entity in _step_field_sources(step):
            for field, entry in entity.items():
                if not isinstance(entry, dict):
                    continue
                row = {
                    "item": str(field),
                    "document": entry.get("document_title") or "",
                    "page": _page_label(entry),
                    "quote": (entry.get("quote") or "").strip(),
                    "status": _field_status(entry) if entry.get("quote") else "no source found",
                }
                key = tuple(row.values())
                if key not in seen:
                    seen.add(key)
                    rows.append(row)
    for src in status.get("retrieved_sources") or []:
        if not isinstance(src, dict):
            continue
        row = {
            "item": "Retrieved passage",
            "document": src.get("document_title") or src.get("source_name") or "",
            "page": _page_label(src),
            "quote": (src.get("content_preview") or "").strip()[:300],
            "status": "retrieved from knowledge base",
        }
        key = tuple(row.values())
        if key not in seen:
            seen.add(key)
            rows.append(row)
    return rows


def _md_cell(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).replace("|", "\\|").strip()


def sources_markdown(rows: list[dict]) -> str:
    """A "Sources" section as a Markdown table, or "" when there are none."""
    if not rows:
        return ""
    lines = ["## Sources", "", "| " + " | ".join(SOURCE_HEADERS) + " |", "|" + "---|" * len(SOURCE_HEADERS)]
    for r in rows:
        lines.append("| " + " | ".join(_md_cell(r[k]) for k in ("item", "document", "page", "quote", "status")) + " |")
    return "\n".join(lines) + "\n"


def append_sources_to_docx(docx_bytes: bytes, rows: list[dict]) -> bytes:
    """Add a Sources heading and table to the end of a Word document."""
    if not rows:
        return docx_bytes
    import io

    from docx import Document

    from app.services.docx_service import _apply_docx_table_theme

    doc = Document(io.BytesIO(docx_bytes))
    doc.add_heading("Sources", level=2)
    table = doc.add_table(rows=1, cols=len(SOURCE_HEADERS))
    for i, h in enumerate(SOURCE_HEADERS):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, k in enumerate(("item", "document", "page", "quote", "status")):
            cells[i].text = str(r[k] or "")
    _apply_docx_table_theme(table)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def field_source_cell(entry: Any) -> str:
    """One field's source as a single CSV cell: "doc, p. 3: "quote"" or "no source"."""
    if not isinstance(entry, dict) or not entry.get("quote"):
        return "no source"
    where = ", ".join(x for x in (entry.get("document_title") or "", _page_label(entry)) if x)
    quote = re.sub(r"\s+", " ", entry["quote"]).strip()
    return f'{where}: "{quote}"' if where else f'"{quote}"'


# --- quotes a chat-saved document cites ------------------------------------

def check_quote(quote: str, document_text: str | None) -> str:
    """Whether *quote* appears in *document_text*, as a Sources status.

    Uses the extraction matcher, which folds smart quotes, dashes, ligatures
    and whitespace the way PDF text layers differ from model output."""
    if document_text is None:
        return "not checked"
    if not (quote or "").strip():
        return "no quote given"
    from app.services.extraction_sources import find_quote_offset

    return "quote found" if find_quote_offset(document_text, quote) is not None else "quote not found in the document"
