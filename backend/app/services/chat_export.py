"""Deterministic exports. No model calls, external fetches, storage or publication."""

import io
import re
from html import escape

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas.chat_export import ChatExportRequest

MIME_TYPES = {
    "html": "text/html",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _message_html(message):
    # Raw HTML is displayed as text. Images never cause external requests.
    md = MarkdownIt("commonmark", {"html": False}).enable("table")
    md.renderer.rules["image"] = lambda tokens, idx, *_: escape(tokens[idx].content)
    body = md.render(message.content)
    tools = "".join(f'<p class="tool">{escape(t)}</p>' for t in message.tools)
    evidence = f'<pre class="evidence">{escape(message.evidence)}</pre>' if message.evidence else ""
    return f'<section><h2>{message.role.title()}</h2>{body}{tools}{evidence}</section>'


def _inline(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    for child in node.children:
        _inline(paragraph, child, bold or node.name in ("strong", "b"), italic or node.name in ("em", "i"))
    if node.name == "a" and node.get("href"):
        # Keep the actual destination readable even when printed.
        paragraph.add_run(f' ({node["href"]})')


def _docx(soup):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Chat export", 0)

    def blocks(parent):
        for node in parent.children:
            if not isinstance(node, Tag):
                continue
            if node.name == "table":
                rows = node.select("tr")
                cols = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Light Shading Accent 1"
                for row in rows:
                    cells = table.add_row().cells
                    for cell, source in zip(cells, row.find_all(["th", "td"], recursive=False)):
                        _inline(cell.paragraphs[0], source, bold=source.name == "th")
            elif node.name in ("section", "ul", "ol", "blockquote"):
                blocks(node)
            elif node.name == "li":
                p = doc.add_paragraph(style="List Number" if node.parent.name == "ol" else "List Bullet")
                for child in node.children:
                    if isinstance(child, Tag) and child.name in ("ul", "ol"):
                        blocks(child)
                    else:
                        _inline(p, child)
            elif node.name == "pre":
                p = doc.add_paragraph(node.get_text())
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
            elif node.name == "hr":
                doc.add_paragraph("—")
            else:
                level = int(node.name[1]) if re.fullmatch(r"h[1-6]", node.name) else 0
                p = doc.add_heading(level=level) if level else doc.add_paragraph()
                _inline(p, node)

    blocks(soup)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _cell(sheet, row, column, value, numeric=False):
    # Force strings to literal cells: formulas and Excel error tokens remain text.
    cell = sheet.cell(row, column)
    cell.value = value
    cell.data_type = "s"
    if numeric and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", value) and len(re.sub(r"\D", "", value)) <= 15:
        cell.value = float(value) if "." in value else int(value)
        cell.data_type = "n"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _xlsx(body, soup):
    wb = Workbook()
    transcript = wb.active
    transcript.title = "Conversation"
    transcript.append(["Message", "Role", "Part", "Content", "Tool summaries", "Source references and warnings"])
    for index, message in enumerate(body.messages, 1):
        values = [message.content, "\n\n".join(message.tools), message.evidence]
        # Excel silently truncates long cells; split them explicitly and retain order.
        chunks = [[v[i:i + 30000] for i in range(0, len(v), 30000)] or [""] for v in values]
        for part in range(max(map(len, chunks))):
            row = transcript.max_row + 1
            transcript.append([index, message.role.title(), part + 1])
            for col, pieces in enumerate(chunks, 4):
                _cell(transcript, row, col, pieces[part] if part < len(pieces) else "")
    for index, table in enumerate(soup.select("table"), 1):
        sheet = wb.create_sheet(f"Table {index}")
        message_index = len(list(table.find_parent("section").find_previous_siblings("section"))) + 1
        sheet.append([f"From message {message_index}; source references are on the Conversation sheet."])
        for row_index, row in enumerate(table.select("tr"), 2):
            for col, source in enumerate(row.find_all(["th", "td"], recursive=False), 1):
                value = source.get_text()
                if len(value) > 30000:
                    raise ValueError("A table cell is too long for Excel. Use HTML or Word for this export.")
                _cell(sheet, row_index, col, value, numeric=source.name == "td")
        sheet.freeze_panes = "A3"
    for sheet in wb:
        header = 1 if sheet == transcript else 2
        for cell in sheet[header]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 42
        sheet.freeze_panes = "A2" if sheet == transcript else "A3"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_chat_export(body: ChatExportRequest) -> bytes:
    content = "".join(_message_html(m) for m in body.messages)
    soup = BeautifulSoup(content, "html.parser")
    tables = soup.select("table")
    if len(tables) > 100 or len(soup.select("td, th")) > 20000:
        raise ValueError("Export has too many table cells. Export fewer messages.")
    if any(len(row.find_all(["th", "td"], recursive=False)) > 100 for row in soup.select("tr")):
        raise ValueError("Export tables support at most 100 columns.")
    if body.format == "docx":
        return _docx(soup)
    if body.format == "xlsx":
        return _xlsx(body, soup)
    page = f'''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>Chat export</title><style>
body{{font:16px/1.6 system-ui,sans-serif;color:#1e293b;max-width:900px;margin:auto;padding:24px}}
section{{border-top:1px solid #cbd5e1;padding:12px 0;overflow-wrap:anywhere}}
table{{display:block;overflow-x:auto;border-collapse:collapse;max-width:100%}}
th,td{{border:1px solid #cbd5e1;padding:8px;text-align:left}}th{{background:#f1f5f9}}
pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:#f8fafc;padding:12px}}
.evidence,.tool{{font-size:13px;color:#475569;white-space:pre-wrap}}h1,h2,h3{{line-height:1.3}}
@media print{{body{{padding:0}}table{{display:table}}}}
</style></head><body><h1>Chat export</h1>{content}</body></html>'''
    return page.encode("utf-8")
