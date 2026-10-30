import io
from unittest.mock import AsyncMock, patch

import pytest
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas.chat_export import ChatExportRequest
from app.services.chat_export import MIME_TYPES, render_chat_export
from tests.test_chat_routes import _auth, _make_user


CONTENT = '''# Award review

**Confirmed** and *pending*: [source](https://example.org/award).

| Item | Amount | Identifier |
| --- | --- | --- |
| Grant | 1250.50 | 00123 |
| =1+1 | -42 | 12345678901234567 |

```html
<script>alert('code example')</script>
```

| Status | Value |
| --- | --- |
| #N/A | @SUM(A1) |
'''


def request(fmt, content=CONTENT):
    return ChatExportRequest(format=fmt, messages=[
        {"role": "user", "content": "Review this award"},
        {"role": "assistant", "content": content, "tools": ["Read two documents"],
         "evidence": "Response stopped; this is partial output.\n[1] Award.pdf · p. 3\nURL: https://example.org/award"},
    ])


def test_html_is_standalone_and_inert_and_keeps_evidence():
    data = render_chat_export(request("html", CONTENT + '\n<script>alert(1)</script>\n\n![image](https://example.org/tracker)\n\n[bad](javascript:alert(1))')).decode()
    soup = BeautifulSoup(data, "html.parser")
    assert not soup.select("script, img, iframe, object")
    assert not any(a.get("href", "").startswith("javascript:") for a in soup.select("a"))
    assert "default-src 'none'" in soup.select_one('meta[http-equiv="Content-Security-Policy"]')["content"]
    assert len(soup.select("table")) == 2
    assert "<script>alert('code example')</script>" in soup.get_text()
    assert "Award.pdf · p. 3" in soup.get_text()
    assert "partial output" in soup.get_text()


def test_word_roundtrip_preserves_headings_tables_code_links_and_sources():
    doc = Document(io.BytesIO(render_chat_export(request("docx"))))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Award review" in text
    assert "https://example.org/award" in text
    assert "<script>alert('code example')</script>" in text
    assert "Read two documents" in text and "Award.pdf · p. 3" in text
    assert any(run.bold for p in doc.paragraphs for run in p.runs if run.text == "Confirmed")
    assert len(doc.tables) == 2
    assert doc.tables[0].rows[1].cells[1].text == "1250.50"


def test_excel_has_all_tables_numeric_cells_and_literal_formulas_and_identifiers():
    wb = load_workbook(io.BytesIO(render_chat_export(request("xlsx"))))
    assert wb.sheetnames == ["Conversation", "Table 1", "Table 2"]
    assert wb["Conversation"]["D3"].value == CONTENT
    assert "Award.pdf · p. 3" in wb["Conversation"]["F3"].value
    assert "message 2" in wb["Table 1"]["A1"].value
    assert wb["Table 1"]["B3"].value == 1250.5
    assert wb["Table 1"]["B4"].value == -42
    assert wb["Table 1"]["C3"].value == "00123"
    assert wb["Table 1"]["C4"].value == "12345678901234567"
    for sheet, address, value in [("Table 1", "A4", "=1+1"), ("Table 2", "A3", "#N/A"), ("Table 2", "B3", "@SUM(A1)")]:
        assert wb[sheet][address].value == value
        assert wb[sheet][address].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in wb for row in sheet for cell in row)


def test_long_transcripts_are_split_without_truncation_and_no_tables_is_valid():
    content = "=literal " + "é" * 70000
    wb = load_workbook(io.BytesIO(render_chat_export(request("xlsx", content))))
    assert wb.sheetnames == ["Conversation"]
    rows = list(wb.active.values)[2:]
    assert "".join(row[3] or "" for row in rows) == content
    assert [row[2] for row in rows] == [1, 2, 3]
    assert all(row[0] == 2 for row in rows)


@pytest.mark.parametrize("content", ["x" * 500001, "bad\x00text", "bad\ufffftext"])
def test_export_size_and_control_characters_are_rejected(content):
    with pytest.raises(ValidationError):
        request("docx", content)


def test_oversized_excel_table_cell_is_explicitly_rejected():
    with pytest.raises(ValueError, match="too long for Excel"):
        render_chat_export(request("xlsx", "|A|B|\n|-|-|\n|" + "x" * 30001 + "|ok|"))


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["html", "docx", "xlsx"])
async def test_authenticated_export_download(client, fmt):
    cookies, headers = _auth()
    with patch("app.dependencies.decode_token", return_value={"sub": "testuser", "type": "access"}), patch("app.dependencies.User") as users:
        users.find_one = AsyncMock(return_value=_make_user())
        response = await client.post("/api/chat/export", json=request(fmt).model_dump(), cookies=cookies, headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(MIME_TYPES[fmt])
    assert response.headers["cache-control"] == "no-store"
    assert f'chat-export.{fmt}' in response.headers["content-disposition"]
    assert len(response.content) > 500


@pytest.mark.asyncio
async def test_export_requires_authentication(client):
    response = await client.post("/api/chat/export", json=request("html").model_dump())
    assert response.status_code in (401, 403)
