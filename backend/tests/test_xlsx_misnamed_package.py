"""An .xlsx that is not an Excel workbook gets a message saying what it is.

openpyxl refuses a zip whose [Content_Types].xml declares no workbook
("File contains no valid workbook part"), and the MarkItDown fallback reads
.xlsx through openpyxl too, so the user saw "File conversion failed after 1
attempts: - XlsxConverter threw OSError …" and Sentry got an error per upload
(Sentry 7598762528).
"""

import zipfile

import openpyxl
import pytest

from app.services.document_readers import (
    DocumentReadError,
    extract_text_from_file,
    extract_text_from_xlsx,
)


def _package(path, content_type, extra=None):
    types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/main.bin" ContentType="{content_type}"/>'
        "</Types>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", types)
        for name, data in (extra or {}).items():
            zf.writestr(name, data)
    return str(path)


@pytest.mark.parametrize("content_type,expected", [
    ("application/vnd.ms-excel.sheet.binary.macroEnabled.main", "Excel Binary Workbook (.xlsb)"),
    ("application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml", "a Word document"),
    ("application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml", "a PowerPoint presentation"),
])
def test_a_misnamed_office_package_is_named(tmp_path, content_type, expected):
    path = _package(tmp_path / "report.xlsx", content_type)
    with pytest.raises(DocumentReadError) as raised:
        extract_text_from_xlsx(path)
    assert expected in str(raised.value)
    assert "XlsxConverter" not in str(raised.value)


def test_an_opendocument_spreadsheet_is_named(tmp_path):
    path = tmp_path / "budget.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", "<office:document-content/>")
    with pytest.raises(DocumentReadError, match=r"OpenDocument spreadsheet \(\.ods\)"):
        extract_text_from_xlsx(str(path))


def test_a_package_declaring_no_workbook_says_so(tmp_path):
    path = _package(tmp_path / "export.xlsx", "application/x-something-else")
    with pytest.raises(DocumentReadError, match="contains no Excel workbook"):
        extract_text_from_xlsx(path)


def test_the_reported_files_styles_only_package_says_so(tmp_path):
    """The exact manifest of the file behind Sentry 7764138486/7598762528:
    styles, theme and document properties, and no workbook or sheet at all —
    a damaged .xlsx, not a renamed format, so none of the named formats match."""
    path = tmp_path / "5751456A.xlsx"
    types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        "</Types>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", types)
        zf.writestr("xl/styles.xml", "<styleSheet/>")
    with pytest.raises(DocumentReadError, match="contains no Excel workbook"):
        extract_text_from_xlsx(str(path))


def test_the_generic_reader_passes_the_message_through_unwrapped(tmp_path):
    """extract_text_from_file re-raises a DocumentReadError as-is: no second
    "Could not read this xlsx:" prefix, and no error-level log for Sentry."""
    path = _package(tmp_path / "report.xlsx", "application/vnd.ms-excel.sheet.binary.macroEnabled.main")
    with pytest.raises(DocumentReadError) as raised:
        extract_text_from_file(path, "xlsx")
    assert str(raised.value).startswith("This file is an Excel Binary Workbook")


def test_a_real_workbook_still_reads(tmp_path):
    path = tmp_path / "ok.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Award amount"
    wb.active["B1"] = 50000
    wb.save(path)
    text = extract_text_from_xlsx(str(path))
    assert "Award amount" in text and "50000" in text


# Review of #982: the sniffer ran on every openpyxl failure and matched
# Default extension entries, so a real workbook with an embedded Word object
# that openpyxl failed to parse for an unrelated reason was called "a Word
# document" instead of falling back to MarkItDown.


def _with_embedded_docx_default(path):
    """Rewrite a real .xlsx so its manifest declares an embedded .docx."""
    import shutil

    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"<Default ",
                    b'<Default Extension="docx" ContentType="application/vnd.openxmlformats-'
                    b'officedocument.wordprocessingml.document"/><Default ',
                    1,
                )
            dst.writestr(item, data)
    shutil.move(tmp, path)


def test_a_real_workbook_that_fails_to_parse_still_falls_back(tmp_path, monkeypatch):
    path = tmp_path / "budget.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Award amount"
    wb.save(path)
    _with_embedded_docx_default(path)

    def broken_load(*_a, **_k):
        raise TypeError("expected <class 'openpyxl.styles.fills.Fill'>")

    monkeypatch.setattr(openpyxl, "load_workbook", broken_load)
    monkeypatch.setattr(
        "app.services.document_readers.convert_to_markdown",
        lambda *_a, **_k: "FALLBACK",
    )
    assert extract_text_from_xlsx(str(path)) == "FALLBACK"


def test_a_presentation_embedding_a_word_file_is_called_a_presentation(tmp_path):
    path = tmp_path / "deck.xlsx"
    types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="docx" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document"/>'
        '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
        "</Types>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", types)
    with pytest.raises(DocumentReadError, match="a PowerPoint presentation"):
        extract_text_from_xlsx(str(path))
