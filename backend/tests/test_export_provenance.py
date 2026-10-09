"""Results keep their evidence when they leave Vandalizer (#1008).

Extraction captures a quote and page per field, and knowledge-base steps
record the passages they retrieved, but no download carried either, and chat
could only save Markdown or plain text. A colleague or PI handed the file had
no way to check it.
"""

import csv
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from openpyxl import load_workbook

from app.routers.workflows import _render_workflow_output
from app.services import chat_tools
from app.services.export_provenance import check_quote, collect_run_sources, sources_markdown

FIELD_SOURCES = [{
    "Award Number": {"quote": "Award No. 2401234", "page": 1, "document_title": "noa.pdf", "verified": True, "support": "supported"},
    "Direct Costs": {"quote": "", "page": None, "document_title": "noa.pdf", "verified": False, "support": "unverified"},
}]
RETRIEVED = [{"document_title": "2 CFR 200.1", "page": 4, "content_preview": "Equipment means tangible personal property…"}]


def _status(output, *, field_sources=FIELD_SOURCES, retrieved=RETRIEVED):
    return {
        "workflow_name": "NOA Summary",
        "final_output": {"output": output},
        "steps_output": {"Extract": {"output": output, "field_sources": field_sources}},
        "output_step_names": ["Extract"],
        "retrieved_sources": retrieved,
    }


ROWS = [{"Award Number": "2401234", "Direct Costs": "$180,000"}]


class TestRunSources:
    def test_collects_field_quotes_and_retrieved_passages(self):
        rows = collect_run_sources(_status(ROWS))
        assert rows[0] == {"item": "Award Number", "document": "noa.pdf", "page": "p. 1",
                           "quote": "Award No. 2401234", "status": "quote found; value matches it"}
        assert rows[1]["item"] == "Direct Costs" and rows[1]["status"] == "no source found"
        assert rows[2]["item"] == "Retrieved passage" and rows[2]["page"] == "p. 4"

    def test_markdown_section_escapes_pipes(self):
        md = sources_markdown([{"item": "a|b", "document": "d", "page": "", "quote": "q", "status": "s"}])
        assert md.startswith("## Sources") and "a\\|b" in md


class TestWorkflowExports:
    def test_csv_puts_each_fields_source_beside_it(self):
        body, *_ = _render_workflow_output(_status(ROWS), "csv", False)
        table = list(csv.reader(io.StringIO(body.decode())))
        assert table[0] == ["Award Number", "Award Number (source)", "Direct Costs", "Direct Costs (source)"]
        assert table[1][1] == 'noa.pdf, p. 1: "Award No. 2401234"'
        assert table[1][3] == "no source"

    def test_csv_of_a_single_record_gets_a_source_column(self):
        body, *_ = _render_workflow_output(_status(ROWS[0]), "csv", False)
        table = list(csv.reader(io.StringIO(body.decode())))
        assert table[0] == ["Field", "Value", "Source"]
        assert table[1] == ["Award Number", "2401234", 'noa.pdf, p. 1: "Award No. 2401234"']

    def test_csv_without_sources_is_unchanged(self):
        body, *_ = _render_workflow_output(_status(ROWS, field_sources=[], retrieved=[]), "csv", False)
        assert next(csv.reader(io.StringIO(body.decode()))) == ["Award Number", "Direct Costs"]

    def test_markdown_ends_with_the_sources(self):
        body, *_ = _render_workflow_output(_status(ROWS), "markdown", False)
        text = body.decode()
        assert "## Sources" in text and "Award No. 2401234" in text and "2 CFR 200.1" in text

    def test_word_has_a_sources_table(self):
        body, *_ = _render_workflow_output(_status(ROWS), "docx", False)
        doc = Document(io.BytesIO(body))
        assert "Sources" in [p.text for p in doc.paragraphs]
        assert doc.tables[-1].rows[0].cells[0].text == "Item"

    def test_pdf_renders_with_sources(self):
        body, media, ext, _ = _render_workflow_output(_status(ROWS), "pdf", False)
        assert body.startswith(b"%PDF") and ext == "pdf"


@pytest.mark.parametrize("quote,text,expected", [
    ("The maximum direct-cost request is $180,000.", "3. Budget\nThe maximum direct-cost request is $180,000. Indirect…", "quote found"),
    ("The maximum direct‑cost request is $180,000.", "The maximum direct-cost request is $180,000.", "quote found"),
    ("The cap is $150,000.", "The maximum direct-cost request is $180,000.", "quote not found in the document"),
    ("anything", None, "not checked"),
    ("", "text", "no quote given"),
])
def test_check_quote(quote, text, expected):
    assert check_quote(quote, text) == expected


class TestChatSavesDeliverables:
    def _ctx(self):
        ctx = MagicMock()
        ctx.deps.user_id, ctx.deps.team_id = "ra-1", "team-1"
        return ctx

    @pytest.mark.asyncio
    async def test_cited_quotes_are_checked_against_documents_the_user_can_read(self):
        docs = {
            "mine": SimpleNamespace(user_id="ra-1", team_id=None, raw_text="The maximum direct-cost request is $180,000."),
            "other": SimpleNamespace(user_id="someone", team_id="other-team", raw_text="secret"),
        }
        # Each document is read once, in the order first cited.
        with patch.object(chat_tools, "SmartDocument") as Doc:
            Doc.find_one = AsyncMock(side_effect=[docs["mine"], docs["other"]])
            rows = await chat_tools._checked_save_sources([
                {"document_uuid": "mine", "document_title": "notice.pdf", "page": 3, "quote": "maximum direct-cost request is $180,000", "item": "Cap"},
                {"document_uuid": "mine", "document_title": "notice.pdf", "page": 3, "quote": "cap is $150,000", "item": "Cap"},
                {"document_uuid": "other", "document_title": "x.pdf", "quote": "secret", "item": "X"},
                {"document_title": "no uuid", "quote": "q", "item": "Y"},
            ], self._ctx())
        assert [r["status"] for r in rows] == [
            "quote found", "quote not found in the document", "not checked", "not checked",
        ]
        assert rows[0]["page"] == "p. 3"

    def test_a_spreadsheet_holds_the_table_and_a_sources_sheet(self):
        rows = [{"item": "Cap", "document": "notice.pdf", "page": "p. 3", "quote": "q", "status": "quote found"}]
        data = chat_tools._render_saved_file(
            "xlsx", "Budget check.xlsx",
            "Intro\n\n| Item | Draft | Cap |\n|---|---|---|\n| Direct costs | $200,000 | $180,000 |\n", rows,
        )
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Budget check", "Sources"]
        assert [[c.value for c in r] for r in wb["Budget check"].iter_rows()] == [
            ["Item", "Draft", "Cap"], ["Direct costs", "$200,000", "$180,000"],
        ]
        assert [c.value for c in wb["Sources"][2]] == ["Cap", "notice.pdf", "p. 3", "q", "quote found"]

    def test_a_word_memo_ends_with_its_sources(self):
        rows = [{"item": "Cap", "document": "notice.pdf", "page": "p. 3", "quote": "q", "status": "quote found"}]
        doc = Document(io.BytesIO(chat_tools._render_saved_file("docx", "Memo.docx", "# Memo to PI\n\nOver the cap.", rows)))
        texts = [p.text for p in doc.paragraphs]
        assert texts[0] == "Memo to PI" and "Sources" in texts

    @pytest.mark.asyncio
    async def test_a_spreadsheet_without_a_table_is_refused_with_a_way_forward(self):
        result = await chat_tools.save_to_folder(self._ctx(), title="Notes", content="Just prose.", extension="xlsx")
        assert "Markdown table" in result["error"] and "docx" in result["hint"]

    @pytest.mark.asyncio
    async def test_unknown_types_name_the_four_that_work(self):
        result = await chat_tools.save_to_folder(self._ctx(), title="Notes", content="x", extension="pdf")
        assert "'md', 'txt', 'docx' or 'xlsx'" in result["error"]
